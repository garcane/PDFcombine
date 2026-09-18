"""Unit tests for the Word reading and conversion layer.

Run with::

    python -m pytest -q
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from core.docx_converter import DocxConvertTool, OutputFormat, PdfEngine, available_pdf_engines
from core.docx_reader import read_docx
from core.exceptions import PdfReadError, PdfWriteError, ValidationError
from core.markdown_writer import document_to_markdown, documents_to_markdown
from core.utils import discover_docx, is_docx
from core.validator import inspect_docx


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #
def make_docx(path: Path, title: str = "Test Document") -> Path:
    """Create a Word document exercising every structure the reader supports."""
    document = Document()
    document.add_heading(title, level=1)
    paragraph = document.add_paragraph("Plain text with ")
    paragraph.add_run("bold").bold = True
    paragraph.add_run(" and ")
    paragraph.add_run("italic").italic = True
    paragraph.add_run(" runs, plus a * star.")
    document.add_paragraph("Bullet one", style="List Bullet")
    document.add_paragraph("Bullet two", style="List Bullet")
    document.add_paragraph("Step one", style="List Number")
    document.add_paragraph("Step two", style="List Number")
    document.add_heading("Subsection", level=2)
    document.add_paragraph("A quotation.", style="Quote")

    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Key"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Alpha"
    table.rows[1].cells[1].text = "1"

    document.save(path)
    return path


@pytest.fixture()
def sample(tmp_path: Path) -> Path:
    """A single Word document."""
    return make_docx(tmp_path / "sample.docx")


@pytest.fixture()
def tree(tmp_path: Path) -> Path:
    """A folder of Word documents including a sub-folder and noise files."""
    root = tmp_path / "input"
    (root / "Nested").mkdir(parents=True)
    make_docx(root / "b.docx", "Bravo")
    make_docx(root / "a.docx", "Alpha")
    make_docx(root / "Nested" / "c.docx", "Charlie")
    (root / "~$a.docx").write_bytes(b"word lock file")
    (root / "notes.txt").write_text("not a document", encoding="utf-8")
    return root


def infos(paths: list[Path], base: Path | None = None):
    """Inspect every path, returning the metadata objects."""
    return [inspect_docx(path, base_folder=base) for path in paths]


# --------------------------------------------------------------------------- #
# Discovery and validation
# --------------------------------------------------------------------------- #
def test_is_docx_rejects_word_lock_files(tmp_path: Path) -> None:
    assert is_docx(tmp_path / "report.docx")
    assert not is_docx(tmp_path / "~$report.docx")
    assert not is_docx(tmp_path / "report.doc")


def test_discover_docx_is_not_recursive_by_default(tree: Path) -> None:
    found = [path.name for path in discover_docx(tree)]
    assert found == ["a.docx", "b.docx"]


def test_discover_docx_recursive_includes_subfolders(tree: Path) -> None:
    found = [path.name for path in discover_docx(tree, recursive=True)]
    assert found == ["a.docx", "b.docx", "c.docx"]


def test_inspect_docx_records_relative_path(tree: Path) -> None:
    info = inspect_docx(tree / "Nested" / "c.docx", base_folder=tree)
    assert info.relative_path == Path("Nested/c.docx")
    assert info.table_count == 1
    assert info.paragraph_count > 0


def test_inspect_docx_rejects_non_docx(tmp_path: Path) -> None:
    other = tmp_path / "notes.txt"
    other.write_text("hello", encoding="utf-8")
    with pytest.raises(PdfReadError):
        inspect_docx(other)


def test_inspect_docx_rejects_corrupt_file(tmp_path: Path) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"definitely not a zip archive")
    with pytest.raises(PdfReadError):
        inspect_docx(broken)


def test_inspect_docx_rejects_missing_file(tmp_path: Path) -> None:
    with pytest.raises(PdfReadError):
        inspect_docx(tmp_path / "nope.docx")


# --------------------------------------------------------------------------- #
# Reading and Markdown
# --------------------------------------------------------------------------- #
def test_read_docx_captures_structure(sample: Path) -> None:
    document = read_docx(sample)
    kinds = [block.kind for block in document.blocks]
    assert "heading" in kinds
    assert "bullet" in kinds
    assert "number" in kinds
    assert "quote" in kinds
    assert "table" in kinds


def test_markdown_output(sample: Path) -> None:
    markdown = document_to_markdown(read_docx(sample))
    assert markdown.startswith("# Test Document")
    assert "**bold**" in markdown
    assert "*italic*" in markdown
    assert "\\*" in markdown  # literal asterisks are escaped
    assert "- Bullet one\n- Bullet two" in markdown  # tight list
    assert "1. Step one" in markdown and "2. Step two" in markdown
    assert "## Subsection" in markdown
    assert "> A quotation." in markdown
    assert "| Key | Value |" in markdown
    assert "| --- | --- |" in markdown


def test_collated_markdown_separates_documents(tmp_path: Path) -> None:
    first = read_docx(make_docx(tmp_path / "one.docx", "One"))
    second = read_docx(make_docx(tmp_path / "two.docx", "Two"))
    combined = documents_to_markdown([first, second])
    assert combined.count("\n---\n") == 1
    assert "One" in combined and "Two" in combined


# --------------------------------------------------------------------------- #
# Conversion - one file per document
# --------------------------------------------------------------------------- #
def test_convert_each_to_markdown_mirrors_structure(tree: Path, tmp_path: Path) -> None:
    out = tmp_path / "out"
    files = infos(discover_docx(tree, recursive=True), base=tree)

    result = DocxConvertTool().convert(files, out, OutputFormat.MARKDOWN, mirror_structure=True)

    assert result.file_count == 3
    assert (out / "a.md").exists()
    assert (out / "Nested" / "c.md").exists()


def test_convert_each_flat_when_not_mirroring(tree: Path, tmp_path: Path) -> None:
    out = tmp_path / "flat"
    files = infos(discover_docx(tree, recursive=True), base=tree)

    DocxConvertTool().convert(files, out, OutputFormat.MARKDOWN, mirror_structure=False)

    assert sorted(p.name for p in out.glob("*.md")) == ["a.md", "b.md", "c.md"]
    assert not (out / "Nested").exists()


def test_convert_each_to_pdf(tree: Path, tmp_path: Path) -> None:
    out = tmp_path / "pdf"
    files = infos(discover_docx(tree), base=tree)

    result = DocxConvertTool().convert(
        files, out, OutputFormat.PDF, engine=PdfEngine.BUILTIN
    )

    assert result.file_count == 2
    for path in result.output_paths:
        assert path.suffix == ".pdf"
        assert len(PdfReader(str(path)).pages) >= 1


def test_convert_keeps_existing_files_when_not_overwriting(sample: Path, tmp_path: Path) -> None:
    out = tmp_path / "twice"
    files = infos([sample])

    DocxConvertTool().convert(files, out, OutputFormat.MARKDOWN)
    second = DocxConvertTool().convert(files, out, OutputFormat.MARKDOWN)

    assert "(2)" in second.output_paths[0].name


def test_convert_overwrites_when_requested(sample: Path, tmp_path: Path) -> None:
    out = tmp_path / "over"
    files = infos([sample])

    first = DocxConvertTool().convert(files, out, OutputFormat.MARKDOWN)
    second = DocxConvertTool().convert(files, out, OutputFormat.MARKDOWN, overwrite=True)

    assert first.output_paths[0] == second.output_paths[0]
    assert len(list(out.glob("*.md"))) == 1


def test_convert_reports_unreadable_documents_without_aborting(
    sample: Path, tmp_path: Path
) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a docx")
    good = inspect_docx(sample)
    # Bypass validation to simulate a file that breaks during conversion.
    bad = inspect_docx(sample)
    object.__setattr__(bad, "path", broken)

    result = DocxConvertTool().convert([good, bad], tmp_path / "mixed", OutputFormat.MARKDOWN)

    assert result.file_count == 1
    assert len(result.warnings) == 1
    assert "broken.docx" in result.warnings[0]


def test_convert_requires_at_least_one_file(tmp_path: Path) -> None:
    with pytest.raises(ValidationError):
        DocxConvertTool().convert([], tmp_path / "out", OutputFormat.MARKDOWN)


def test_convert_raises_when_nothing_can_be_read(tmp_path: Path) -> None:
    sample = make_docx(tmp_path / "ok.docx")
    info = inspect_docx(sample)
    object.__setattr__(info, "path", tmp_path / "missing.docx")

    with pytest.raises(PdfWriteError):
        DocxConvertTool().convert([info], tmp_path / "out", OutputFormat.MARKDOWN)


# --------------------------------------------------------------------------- #
# Conversion - collated
# --------------------------------------------------------------------------- #
def test_collate_to_single_pdf_preserves_order(tree: Path, tmp_path: Path) -> None:
    out = tmp_path / "combined"
    files = infos(discover_docx(tree, recursive=True), base=tree)

    result = DocxConvertTool().convert(
        files,
        out,
        OutputFormat.PDF,
        combine=True,
        combined_name="All",
        engine=PdfEngine.BUILTIN,
    )

    assert result.file_count == 1
    combined = result.output_paths[0]
    assert combined.name == "All.pdf"
    reader = PdfReader(str(combined))
    assert len(reader.pages) == 3  # one page break per document
    assert "Alpha" in reader.pages[0].extract_text()
    assert "Charlie" in reader.pages[2].extract_text()


def test_collate_to_single_markdown(tree: Path, tmp_path: Path) -> None:
    out = tmp_path / "combined_md"
    files = infos(discover_docx(tree, recursive=True), base=tree)

    result = DocxConvertTool().convert(files, out, OutputFormat.MARKDOWN, combine=True)

    assert result.file_count == 1
    text = result.output_paths[0].read_text(encoding="utf-8")
    assert result.output_paths[0].name == "Combined.md"
    assert text.count("\n---\n") == 2


def test_collated_name_gains_the_right_extension(sample: Path, tmp_path: Path) -> None:
    result = DocxConvertTool().convert(
        infos([sample]), tmp_path / "x", OutputFormat.MARKDOWN, combine=True, combined_name="Book"
    )
    assert result.output_paths[0].name == "Book.md"


def test_single_document_converts_on_its_own(sample: Path, tmp_path: Path) -> None:
    result = DocxConvertTool().convert(
        infos([sample]), tmp_path / "one", OutputFormat.PDF, engine=PdfEngine.BUILTIN
    )
    assert result.file_count == 1
    assert len(PdfReader(str(result.output_paths[0])).pages) == 1


# --------------------------------------------------------------------------- #
# Engines and cancellation
# --------------------------------------------------------------------------- #
def test_builtin_engine_is_always_available() -> None:
    assert PdfEngine.BUILTIN in available_pdf_engines()


def test_cancellation_stops_a_batch(tree: Path, tmp_path: Path) -> None:
    from core.exceptions import OperationCancelled
    from core.progress import CancellationToken

    token = CancellationToken()
    token.cancel()
    files = infos(discover_docx(tree, recursive=True), base=tree)

    with pytest.raises(OperationCancelled):
        DocxConvertTool().convert(
            files, tmp_path / "cancel", OutputFormat.MARKDOWN, token=token
        )
